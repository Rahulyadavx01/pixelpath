import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    BG_COLOR = RGBColor(11, 15, 25)         # #0B0F19 Dark Navy Slate
    CARD_BG = RGBColor(22, 30, 46)          # #161E2E Card Container
    BORDER_COLOR = RGBColor(45, 55, 75)     # #2D374B Subtle border
    TEXT_LIGHT = RGBColor(248, 250, 252)    # #F8FAFC White/light
    TEXT_MUTED = RGBColor(148, 163, 184)    # #94A3B8 Muted slate
    ACCENT_PINK = RGBColor(236, 72, 153)    # #EC4899 Pink Accent
    ACCENT_SKY = RGBColor(56, 189, 248)     # #38BDF8 Sky Blue Accent
    ACCENT_AMBER = RGBColor(245, 158, 11)   # #F59E0B Amber Accent
    ACCENT_PURPLE = RGBColor(168, 85, 247)  # #A855F7 Purple Accent

    blank_layout = prs.slide_layouts[6]

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, tag_text, title_text):
        # Category Tag
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.5), Inches(0.4))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = tag_text.upper()
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = ACCENT_SKY

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.5), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_LIGHT

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)

    # Accent Glow Pill
    pill = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(3.6), Inches(0.45))
    pill.fill.solid()
    pill.fill.fore_color.rgb = CARD_BG
    pill.line.color.rgb = ACCENT_PINK
    pill.line.width = Pt(1.5)
    tf_pill = pill.text_frame
    p_pill = tf_pill.paragraphs[0]
    p_pill.text = "⚡ GAMIFIED LEARNING ECOSYSTEM"
    p_pill.font.size = Pt(11)
    p_pill.font.bold = True
    p_pill.font.color.rgb = ACCENT_PINK
    p_pill.alignment = PP_ALIGN.CENTER

    # Project Title
    t_box = slide1.shapes.add_textbox(Inches(0.8), Inches(2.4), Inches(11.5), Inches(1.4))
    tf = t_box.text_frame
    p = tf.paragraphs[0]
    p.text = "PIXELPATH"
    p.font.size = Pt(56)
    p.font.bold = True
    p.font.color.rgb = TEXT_LIGHT

    # Subtitle
    sub_box = slide1.shapes.add_textbox(Inches(0.8), Inches(3.8), Inches(11.5), Inches(1.0))
    tf_sub = sub_box.text_frame
    tf_sub.word_wrap = True
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "Institutional Gamified Learning & Competitive Coding Platform\nTurning Computer Science Curriculum into an Addictive Daily Habit"
    p_sub.font.size = Pt(20)
    p_sub.font.color.rgb = ACCENT_SKY

    # Meta Details Card
    meta_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.2), Inches(11.7), Inches(1.5))
    meta_card.fill.solid()
    meta_card.fill.fore_color.rgb = CARD_BG
    meta_card.line.color.rgb = BORDER_COLOR
    tf_meta = meta_card.text_frame
    tf_meta.word_wrap = True
    p_m1 = tf_meta.paragraphs[0]
    p_m1.text = "Author: Rahul Yadav  |  Target Domain: Computer Science & Engineering  |  Submission: October 2026"
    p_m1.font.size = Pt(13)
    p_m1.font.bold = True
    p_m1.font.color.rgb = TEXT_LIGHT

    p_m2 = tf_meta.add_paragraph()
    p_m2.text = "Core Components: Real-Time Compiler Sandbox  •  Dynamic Curriculum DAG  •  Mains-XXXXX Auth  •  Faculty Governance"
    p_m2.font.size = Pt(12)
    p_m2.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------
    # SLIDE 2: Problem & Motivation
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(slide2, "Background & Motivation", "The Challenge in Modern Computer Science Education")

    col_data = [
        ("Disjointed Practice Tools", "Students switch across isolated tools for lectures, quizzes, and problem solving, resulting in disjointed learning contexts and high dropout rates.", ACCENT_PINK),
        ("Lack of Habit Loops", "Without continuous feedback and healthy gamification, coding practice becomes cramming right before exams rather than daily muscle memory.", ACCENT_AMBER),
        ("Rigid Curriculum Silos", "Faculty cannot dynamically adapt module prerequisites or track real-time comprehension velocity across diverse student cohorts.", ACCENT_SKY)
    ]

    for i, (head, desc, color) in enumerate(col_data):
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 4.0), Inches(2.0), Inches(3.7), Inches(4.7))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = color
        card.line.width = Pt(1.5)
        
        tf_c = card.text_frame
        tf_c.word_wrap = True
        p1 = tf_c.paragraphs[0]
        p1.text = f"0{i+1}"
        p1.font.size = Pt(18)
        p1.font.bold = True
        p1.font.color.rgb = color

        p2 = tf_c.add_paragraph()
        p2.text = head
        p2.font.size = Pt(18)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_LIGHT
        
        p3 = tf_c.add_paragraph()
        p3.text = "\n" + desc
        p3.font.size = Pt(13)
        p3.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------
    # SLIDE 3: The PIXELPATH Solution
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(slide3, "Core Solution Overview", "PIXELPATH: The Unified Learning Architecture")

    # Left Card: Student Experience
    s_card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(5.6), Inches(4.9))
    s_card.fill.solid()
    s_card.fill.fore_color.rgb = CARD_BG
    s_card.line.color.rgb = ACCENT_SKY
    s_card.line.width = Pt(1.5)
    tf_s = s_card.text_frame
    tf_s.word_wrap = True
    ps1 = tf_s.paragraphs[0]
    ps1.text = "STUDENT PORTAL"
    ps1.font.size = Pt(14)
    ps1.font.bold = True
    ps1.font.color.rgb = ACCENT_SKY

    ps2 = tf_s.add_paragraph()
    ps2.text = "Engaging, Gamified & Fair Mastery"
    ps2.font.size = Pt(18)
    ps2.font.bold = True
    ps2.font.color.rgb = TEXT_LIGHT

    items_s = [
        "Code Arena: Instant execution in C++17, Python 3, JavaScript, Java with hidden test drivers.",
        "Curriculum Progress Map: Visual milestone graph with prerequisite gating.",
        "Verified XP Economy: Strict non-academic XP rules (never replaces exam marks).",
        "Habit Retention: Daily quest sprints, streaks with freeze shields, and cohort leagues."
    ]
    for it in items_s:
        p_it = tf_s.add_paragraph()
        p_it.text = "• " + it
        p_it.font.size = Pt(12)
        p_it.font.color.rgb = TEXT_MUTED

    # Right Card: Faculty Experience
    f_card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.9), Inches(5.7), Inches(4.9))
    f_card.fill.solid()
    f_card.fill.fore_color.rgb = CARD_BG
    f_card.line.color.rgb = ACCENT_PURPLE
    f_card.line.width = Pt(1.5)
    tf_f = f_card.text_frame
    tf_f.word_wrap = True
    pf1 = tf_f.paragraphs[0]
    pf1.text = "FACULTY PORTAL"
    pf1.font.size = Pt(14)
    pf1.font.bold = True
    pf1.font.color.rgb = ACCENT_PURPLE

    pf2 = tf_f.add_paragraph()
    pf2.text = "Academic Oversight & Curriculum Engine"
    pf2.font.size = Pt(18)
    pf2.font.bold = True
    pf2.font.color.rgb = TEXT_LIGHT

    items_f = [
        "Dynamic Curriculum Builder: Create courses, modules, and lessons without hardcoded UI limits.",
        "Prerequisite DAG Validation: Circular dependency blocking & affected student impact analysis.",
        "Daily Quiz Studio: Author questions, randomize pools, and configure delivery windows.",
        "Cohort Analytics: Real-time tracking of student completion velocity and code performance."
    ]
    for it in items_f:
        p_it = tf_f.add_paragraph()
        p_it.text = "• " + it
        p_it.font.size = Pt(12)
        p_it.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------
    # SLIDE 4: Code Arena Deep Dive
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(slide4, "Feature Deep Dive", "Real-Time Sandboxed Code Arena")

    box_data = [
        ("Multi-Language Support", "Real local execution for C++17 (-O2), Python 3, JavaScript (Node), and Java. No mock results.", ACCENT_PINK),
        ("Hidden Driver Execution", "Automatically wraps solution classes in hidden main drivers to parse standard input and format output.", ACCENT_SKY),
        ("Safety & Resource Limits", "Sandboxed execution with 2.0s timeout, memory limit protection, and clean process isolation.", ACCENT_AMBER),
        ("Output Normalization", "Intelligent whitespace trimming and order-agnostic matching for multi-solution algorithms like Two Sum.", ACCENT_PURPLE)
    ]

    for i, (h, d, col) in enumerate(box_data):
        row = i // 2
        col_idx = i % 2
        card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + col_idx * 5.9), Inches(1.9 + row * 2.5), Inches(5.6), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = col
        card.line.width = Pt(1.5)

        tf_c = card.text_frame
        tf_c.word_wrap = True
        p1 = tf_c.paragraphs[0]
        p1.text = h
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = col

        p2 = tf_c.add_paragraph()
        p2.text = "\n" + d
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------
    # SLIDE 5: Dynamic Course Curriculum Engine
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(slide5, "Curriculum Architecture", "Fully Data-Driven Prerequisite Engine")

    card_c = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(11.7), Inches(4.9))
    card_c.fill.solid()
    card_c.fill.fore_color.rgb = CARD_BG
    card_c.line.color.rgb = BORDER_COLOR
    tf_cc = card_c.text_frame
    tf_cc.word_wrap = True

    p = tf_cc.paragraphs[0]
    p.text = "How PIXELPATH Manages Academic Progression:"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_SKY

    points = [
        ("Zero Hardcoded Data", "All courses, modules, lessons, and XP rewards are queried dynamically via API endpoints and database tables."),
        ("Prerequisite Graph (DAG)", "Modules unlock strictly when all prerequisites are fulfilled. Circular dependencies (A -> B -> A) are blocked automatically."),
        ("Server-Enforced Gating", "Students cannot bypass UI locks by hitting the API directly; complete-module endpoints verify prerequisites server-side."),
        ("Non-Destructive Evolution", "If faculty update a published module while a student has progress, completed status and earned XP are preserved."),
        ("Audit Logging", "Every syllabus alteration, module edit, and deletion is recorded in a tamper-proof audit log for academic governance.")
    ]

    for title, desc in points:
        p_t = tf_cc.add_paragraph()
        p_t.text = "• " + title + ": " + desc
        p_t.font.size = Pt(12)
        p_t.font.color.rgb = TEXT_LIGHT

    # -------------------------------------------------------------
    # SLIDE 6: Authentication & Security
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6)
    add_header(slide6, "Institutional Identity", "Strict Student ID Format & Anti-Bruteforce")

    auth_cards = [
        ("Regex ID Format: Mains-XXXXX", "Students authenticate with their official institutional ID matching ^Mains-\\d{5}$ (e.g. Mains-24813).\nNormalized casing & whitespace on both client and server.", ACCENT_PINK),
        ("Roster-Gated Activation", "New students can only activate accounts if their ID exists in the pre-authorized college roster, preventing unauthorized registrations.", ACCENT_SKY),
        ("Anti-Bruteforce Protection", "Strict rate limiting locks out login attempts after consecutive failures per IP and Student ID, safeguarding user accounts.", ACCENT_AMBER)
    ]

    for i, (h, d, col) in enumerate(auth_cards):
        card = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 4.0), Inches(2.0), Inches(3.7), Inches(4.7))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = col
        card.line.width = Pt(1.5)

        tf_c = card.text_frame
        tf_c.word_wrap = True
        p1 = tf_c.paragraphs[0]
        p1.text = h
        p1.font.size = Pt(17)
        p1.font.bold = True
        p1.font.color.rgb = col

        p2 = tf_c.add_paragraph()
        p2.text = "\n" + d
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------
    # SLIDE 7: Technical Stack & Architecture
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7)
    add_header(slide7, "System Architecture", "Modern Full-Stack Engineering")

    tech_boxes = [
        ("Frontend Application", "• React 18 with TypeScript\n• Tailwind CSS for modern aesthetics\n• Framer Motion fluid animations\n• Monaco Code Editor integration", ACCENT_SKY),
        ("Backend Compiler Server", "• Node.js & Express 5\n• Subprocess execution sandbox\n• Gzip HTTP compression\n• SPA wildcard client routing", ACCENT_PINK),
        ("Database & Operations", "• Supabase PostgreSQL\n• Row-Level Security (RLS)\n• High-performance indexed tables\n• LocalStorage cache fallback", ACCENT_PURPLE),
        ("Production & Deployment", "• Cloudflare Tunnel (HTTP/2)\n• Production bundle optimization\n• 7-Day browser asset caching\n• Automated GitHub CI/CD ready", ACCENT_AMBER)
    ]

    for i, (h, d, col) in enumerate(tech_boxes):
        row = i // 2
        col_idx = i % 2
        card = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + col_idx * 5.9), Inches(1.9 + row * 2.5), Inches(5.6), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = col
        card.line.width = Pt(1.5)

        tf_c = card.text_frame
        tf_c.word_wrap = True
        p1 = tf_c.paragraphs[0]
        p1.text = h
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = col

        p2 = tf_c.add_paragraph()
        p2.text = "\n" + d
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------
    # SLIDE 8: Project Impact & Deliverables
    # -------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8)
    add_header(slide8, "Demonstration & Access", "Live Project Links & Credentials")

    # Links Card
    link_card = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(11.7), Inches(4.9))
    link_card.fill.solid()
    link_card.fill.fore_color.rgb = CARD_BG
    link_card.line.color.rgb = ACCENT_SKY
    link_card.line.width = Pt(1.5)
    tf_l = link_card.text_frame
    tf_l.word_wrap = True

    p_l1 = tf_l.paragraphs[0]
    p_l1.text = "Repository & Live Demonstration Links:"
    p_l1.font.size = Pt(18)
    p_l1.font.bold = True
    p_l1.font.color.rgb = ACCENT_SKY

    p_g = tf_l.add_paragraph()
    p_g.text = "• GitHub Public Repository: https://github.com/Rahulyadavx01/questify"
    p_g.font.size = Pt(13)
    p_g.font.color.rgb = TEXT_LIGHT

    p_d = tf_l.add_paragraph()
    p_d.text = "• Live Application URL: https://thinks-strings-prayers-economics.trycloudflare.com"
    p_d.font.size = Pt(13)
    p_d.font.color.rgb = TEXT_LIGHT

    p_s = tf_l.add_paragraph()
    p_s.text = "\nDemo Credentials for Evaluation:"
    p_s.font.size = Pt(15)
    p_s.font.bold = True
    p_s.font.color.rgb = ACCENT_AMBER

    p_s1 = tf_l.add_paragraph()
    p_s1.text = "• Student Portal: Login with Student ID 'Mains-24813' and password 'password123'"
    p_s1.font.size = Pt(13)
    p_s1.font.color.rgb = TEXT_MUTED

    p_s2 = tf_l.add_paragraph()
    p_s2.text = "• Faculty Portal: Login with Email 'faculty@questify.edu' and password 'admin'"
    p_s2.font.size = Pt(13)
    p_s2.font.color.rgb = TEXT_MUTED

    p_s3 = tf_l.add_paragraph()
    p_s3.text = "\nSummary: PIXELPATH successfully solves the challenge of student habit-building and curriculum governance with an enterprise-grade, high-performance platform."
    p_s3.font.size = Pt(13)
    p_s3.font.bold = True
    p_s3.font.color.rgb = ACCENT_PINK

    # Save presentation
    output_path = "/Users/rahulyadav/Downloads/class challenge/PIXELPATH_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation successfully created at {output_path}")

if __name__ == "__main__":
    create_presentation()
